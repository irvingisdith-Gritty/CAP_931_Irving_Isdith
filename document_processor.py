# CAP 931 - Document Processing Module
# Supports PDF, DOCX, and TXT files

import io
import pymupdf
from docx import Document


def extract_document_text(uploaded_file):
    """
    Extract text from uploaded PDF, DOCX, or TXT files.
    """

    if uploaded_file is None:
        return ""

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    # Limit uploads to 5 MB
    if len(file_bytes) > 5 * 1024 * 1024:
        raise ValueError("Document exceeds the 5 MB limit.")

    if file_name.endswith(".pdf"):

        with pymupdf.open(
            stream=file_bytes,
            filetype="pdf"
        ) as document:

            text = "\n".join(
                page.get_text()
                for page in document
            )

    elif file_name.endswith(".docx"):

        document = Document(io.BytesIO(file_bytes))

        paragraphs = [
            paragraph.text
            for paragraph in document.paragraphs
        ]

        # Include text from Word tables
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    paragraphs.append(cell.text)

        text = "\n".join(paragraphs)

    elif file_name.endswith(".txt"):

        text = file_bytes.decode("utf-8-sig")

    else:
        raise ValueError("Unsupported document format.")

    if not text.strip():
        raise ValueError(
            "No readable text found in the uploaded document."
        )

    # Limit document content sent to the AI
    return text[:15000]